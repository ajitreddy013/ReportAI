import os
import uuid
import time
from typing import Dict, List, Any, Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

class DocumentBuilder:
    """Builds clean, professionally styled Word documents (.docx) from structured report JSON"""

    def __init__(self, outputs_dir: Optional[str] = None):
        if outputs_dir is None:
            backend_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            self.outputs_dir = os.path.join(backend_dir, "outputs")
        else:
            self.outputs_dir = outputs_dir
        os.makedirs(self.outputs_dir, exist_ok=True)

    def create_docx_report(self, report_data: Dict[str, Any]) -> str:
        """
        Builds a full .docx academic report from report JSON
        Returns the filename created inside outputs_dir
        """
        doc = Document()

        # Set page margins (1 inch all around)
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        topic = report_data.get("topic", "Academic Report")
        acad = report_data.get("academic_info", {})
        student_name = acad.get("student_name", "Student Name")
        roll_no = acad.get("roll_no", "N/A")
        college_name = acad.get("college_name", "Sinhgad College of Engineering, Pune")
        department = acad.get("department", "Department of Computer Engineering")
        guide_name = acad.get("guide_name", "Faculty Guide")
        academic_year = acad.get("academic_year", "2025 - 2026")
        format_title = report_data.get("format_title", "Project Report")

        # -------------------------------------------------------------
        # 1. TITLE / COVER PAGE
        # -------------------------------------------------------------
        p_col = doc.add_paragraph()
        p_col.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_col = p_col.add_run(college_name.upper())
        r_col.font.name = "Arial"
        r_col.font.size = Pt(16)
        r_col.font.bold = True
        r_col.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        p_dept = doc.add_paragraph()
        p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_dept = p_dept.add_run(department)
        r_dept.font.name = "Arial"
        r_dept.font.size = Pt(12)
        r_dept.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        doc.add_paragraph()  # spacer
        doc.add_paragraph()

        p_type = doc.add_paragraph()
        p_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_type = p_type.add_run(f"A {format_title.upper()}\nON")
        r_type.font.name = "Arial"
        r_type.font.size = Pt(11)
        r_type.font.bold = True
        r_type.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

        doc.add_paragraph()

        p_top = doc.add_paragraph()
        p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_top = p_top.add_run(f'"{topic}"')
        r_top.font.name = "Arial"
        r_top.font.size = Pt(18)
        r_top.font.bold = True
        r_top.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        doc.add_paragraph()
        doc.add_paragraph()
        doc.add_paragraph()

        # Metadata Table (Submitted by / Guided by)
        table = doc.add_table(rows=1, cols=2)
        table.autofit = True
        table.columns[0].width = Inches(3.0)
        table.columns[1].width = Inches(3.0)

        cell_left = table.cell(0, 0)
        p_sub = cell_left.paragraphs[0]
        r_sub_h = p_sub.add_run("SUBMITTED BY:\n")
        r_sub_h.font.bold = True
        r_sub_h.font.size = Pt(10)
        r_sub_v = p_sub.add_run(f"{student_name}\nRoll No: {roll_no}")
        r_sub_v.font.size = Pt(10)

        cell_right = table.cell(0, 1)
        p_guid = cell_right.paragraphs[0]
        p_guid.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_guid_h = p_guid.add_run("GUIDED BY:\n")
        r_guid_h.font.bold = True
        r_guid_h.font.size = Pt(10)
        r_guid_v = p_guid.add_run(f"{guide_name}\nAcademic Year: {academic_year}")
        r_guid_v.font.size = Pt(10)

        doc.add_page_break()

        # -------------------------------------------------------------
        # 2. TABLE OF CONTENTS SUMMARY
        # -------------------------------------------------------------
        p_toc = doc.add_paragraph()
        r_toc = p_toc.add_run("TABLE OF CONTENTS")
        r_toc.font.name = "Arial"
        r_toc.font.size = Pt(14)
        r_toc.font.bold = True
        r_toc.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        sections = report_data.get("sections", [])
        for sec in sections:
            p_toc_item = doc.add_paragraph()
            p_toc_item.paragraph_format.left_indent = Inches(0.2)
            p_toc_item.paragraph_format.space_after = Pt(4)
            r_num = p_toc_item.add_run(f"{sec.get('number', '')}. {sec.get('title', '')}")
            r_num.font.name = "Arial"
            r_num.font.size = Pt(11)

        doc.add_page_break()

        # -------------------------------------------------------------
        # 3. REPORT BODY SECTIONS
        # -------------------------------------------------------------
        for sec in sections:
            sec_num = sec.get("number", "")
            sec_title = sec.get("title", "")
            sec_content = sec.get("content", "")

            # Section Heading
            p_head = doc.add_paragraph()
            p_head.paragraph_format.space_before = Pt(14)
            p_head.paragraph_format.space_after = Pt(6)
            p_head.paragraph_format.keep_with_next = True
            
            r_h = p_head.add_run(f"{sec_num}. {sec_title}")
            r_h.font.name = "Arial"
            r_h.font.size = Pt(13)
            r_h.font.bold = True
            r_h.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

            # Section Paragraphs
            paragraphs = sec_content.split("\n\n")
            for para in paragraphs:
                para_clean = para.strip()
                if not para_clean:
                    continue

                p_body = doc.add_paragraph()
                p_body.paragraph_format.space_after = Pt(6)
                p_body.paragraph_format.line_spacing = 1.15
                
                # Check for bullet items or subheadings
                if para_clean.startswith("• ") or para_clean.startswith("- "):
                    p_body.paragraph_format.left_indent = Inches(0.25)
                    r_b = p_body.add_run(para_clean)
                    r_b.font.name = "Arial"
                    r_b.font.size = Pt(11)
                elif para_clean.startswith("1.") or para_clean.startswith("2.") or para_clean.startswith("3.") or para_clean.startswith("4."):
                    r_b = p_body.add_run(para_clean)
                    r_b.font.name = "Arial"
                    r_b.font.size = Pt(11)
                else:
                    p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    r_b = p_body.add_run(para_clean)
                    r_b.font.name = "Arial"
                    r_b.font.size = Pt(11)

        # Generate unique filename
        safe_name = "".join(c for c in student_name if c.isalnum() or c in (' ', '_')).rstrip().replace(" ", "_")
        safe_topic = "".join(c for c in topic[:20] if c.isalnum() or c in (' ', '_')).rstrip().replace(" ", "_")
        filename = f"Report_{safe_name}_{safe_topic}_{int(time.time())}.docx"
        file_path = os.path.join(self.outputs_dir, filename)
        doc.save(file_path)
        return filename

document_builder = DocumentBuilder()

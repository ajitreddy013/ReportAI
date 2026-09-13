"""Deep structural analysis of a DOCX: page setup, headers/footers, media,
ordered headings + front matter, tables, fonts. Reusable for reference vs output.

Usage: cd backend && ./venv/bin/python ../analyze_docx_deep.py "<path.docx>"
"""
import sys, os, re, html, zipfile
from collections import Counter, OrderedDict

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "backend"))

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def wtxt(data: bytes) -> str:
    xml = data.decode("utf-8", "ignore")
    return html.unescape("".join(re.findall(r"<w:t(?:\s[^>]*)?>(.*?)</w:t>", xml, re.S)))


def analyze(path: str):
    from docx import Document
    from docx.shared import Pt
    print("=" * 78)
    print("FILE:", path)
    print("SIZE:", f"{os.path.getsize(path):,} bytes")
    print("=" * 78)

    z = zipfile.ZipFile(path)
    names = z.namelist()

    # ---- package parts ----
    media = [n for n in names if n.startswith("word/media/")]
    headers = sorted(n for n in names if re.match(r"word/header\d+\.xml", n))
    footers = sorted(n for n in names if re.match(r"word/footer\d+\.xml", n))
    print("\n-- PACKAGE PARTS --")
    print("media files:", len(media))
    for m in media:
        print("   ", m, f"{z.getinfo(m).file_size:,}b")
    print("header parts:", headers)
    print("footer parts:", footers)

    # ---- section / page setup ----
    doc = Document(path)
    print("\n-- SECTIONS / PAGE SETUP --")
    for i, s in enumerate(doc.sections):
        pg = s.page_width, s.page_height
        mar = (s.left_margin, s.right_margin, s.top_margin, s.bottom_margin)
        def inch(v):
            return round(v.inches, 2) if v is not None else None
        print(f"  section {i}: start_type={s.start_type} orientation={s.orientation}")
        print(f"     page WxH (in): {inch(pg[0])} x {inch(pg[1])}")
        print(f"     margins L/R/T/B (in): {inch(mar[0])}/{inch(mar[1])}/{inch(mar[2])}/{inch(mar[3])}")
        print(f"     header dist={inch(s.header_distance)} footer dist={inch(s.footer_distance)}")
        print(f"     different_first_page={s.different_first_page_header_footer}")
        print(f"     titlePg in xml: {'<w:titlePg' in s._sectPr.xml}")

    # ---- headers / footers text ----
    print("\n-- HEADER TEXT --")
    for h in headers:
        t = wtxt(z.read(h)).strip()
        print(f"  {h}: {t!r}")
    print("-- FOOTER TEXT --")
    for f in footers:
        t = wtxt(z.read(f)).strip()
        print(f"  {f}: {t!r}")

    # ---- ordered paragraph map (styles + headings + front matter) ----
    print("\n-- PARAGRAPH MAP (first 220 non-empty) --")
    paras = doc.paragraphs
    print("total paragraphs:", len(paras))
    shown = 0
    for idx, p in enumerate(paras):
        t = p.text.strip()
        if not t:
            continue
        style = p.style.name if p.style else "?"
        # font of first run
        fsz = None; fname = None; bold = None
        if p.runs:
            r = p.runs[0]
            fsz = r.font.size.pt if r.font.size else None
            fname = r.font.name
            bold = r.font.bold
        flag = ""
        if re.match(r"^\d+(\.\d+)*\s", t) or style.lower().startswith("heading"):
            flag = "  <== HEADING"
        print(f"[{idx:3}] ({style}|{fname}|{fsz}|b={bold}) {t[:90]!r}{flag}")
        shown += 1
        if shown >= 220:
            print("   ... (truncated)")
            break

    # ---- style usage counter ----
    print("\n-- STYLE USAGE --")
    sc = Counter(p.style.name for p in paras if p.style)
    for st, c in sc.most_common(20):
        print(f"   {c:4}  {st}")

    # ---- font usage ----
    print("\n-- FONT USAGE (runs) --")
    fc = Counter()
    szc = Counter()
    for p in paras:
        for r in p.runs:
            if r.font.name:
                fc[r.font.name] += 1
            if r.font.size:
                szc[r.font.size.pt] += 1
    print("  fonts:", dict(fc.most_common(10)))
    print("  sizes:", dict(sorted(szc.items())))

    # ---- tables ----
    print("\n-- TABLES --")
    print("total tables:", len(doc.tables))
    for ti, tb in enumerate(doc.tables):
        rows = len(tb.rows); cols = len(tb.columns)
        first = []
        try:
            for ci in range(cols):
                first.append(tb.cell(0, ci).text.strip()[:30])
        except Exception:
            pass
        print(f"  table {ti}: {rows}x{cols}  row0={first}")
        if rows <= 14:
            for ri in range(rows):
                cells = [tb.cell(ri, ci).text.strip()[:38] for ci in range(cols)]
                print(f"      r{ri}: {cells}")

    z.close()


if __name__ == "__main__":
    analyze(sys.argv[1])

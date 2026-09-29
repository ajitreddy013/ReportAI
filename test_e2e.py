"""
Quick end-to-end test: creates a minimal DOCX template, uploads it,
generates a report, and verifies the output.
"""
import sys
import os
import json

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "backend"))

from docx import Document
from docx.shared import Pt
import tempfile, requests

API = "http://localhost:8000/api"

# ── 1. Create a minimal DOCX template ─────────────────────────────────────────
print("Creating test DOCX template...")
doc = Document()
doc.add_heading("Sample Technical Report Template", 0)
p = doc.add_paragraph("Submitted by: [Student Name]  Roll No: [Roll No]")
p = doc.add_paragraph("Department: [Department]   Guide: [Guide Name]")
p = doc.add_paragraph("Topic: [Topic]   Academic Year: [Academic Year]")
doc.add_heading("Abstract", 1)
doc.add_paragraph("Write your abstract here. Provide a brief overview of the report.")
doc.add_heading("Introduction", 1)
doc.add_paragraph("Write your introduction here. Explain the background and motivation.")
doc.add_heading("Methodology", 1)
doc.add_paragraph("Describe the methodology adopted for this study.")
doc.add_heading("Results and Discussion", 1)
doc.add_paragraph("Present your results here.")
doc.add_heading("Conclusion", 1)
doc.add_paragraph("Summarize the findings and future scope.")
doc.add_heading("References", 1)
doc.add_paragraph("[1] Author, Title, Journal, Year")

tmpfile = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
doc.save(tmpfile.name)
tmpfile.close()
print(f"Template saved: {tmpfile.name}")

# ── 2. Upload ─────────────────────────────────────────────────────────────────
print("\nUploading document...")
with open(tmpfile.name, "rb") as f:
    res = requests.post(f"{API}/upload-doc", files={"file": ("test_template.docx", f)})
assert res.status_code == 200, f"Upload failed: {res.text}"
upload_data = res.json()
print(f"  doc_id={upload_data['doc_id']}")
print(f"  is_template={upload_data['is_template']}")
print(f"  sections={upload_data['headings']}")
print(f"  placeholders={upload_data['placeholders']}")

# ── 3. Generate ───────────────────────────────────────────────────────────────
print("\nGenerating report...")
form_data = {
    "doc_id":         upload_data["doc_id"],
    "mode":           "template",
    "topic":          "Artificial Intelligence in Healthcare",
    "student_name":   "Ajit Reddy",
    "roll_no":        "22CS1001",
    "department":     "Computer Engineering",
    "college_name":   "Sinhgad College of Engineering",
    "guide_name":     "Prof. R. K. Sharma",
    "academic_year":  "2024-25",
    "reference_notes": "Focus on machine learning in diagnostics",
    "images_json":    "[]",
}
res = requests.post(f"{API}/generate", data=form_data)
assert res.status_code == 200, f"Generate failed: {res.text}"
gen_data = res.json()
print(f"  report_id={gen_data['report_id']}")
print(f"  sections={gen_data['sections']}")

# ── 4. Preview ────────────────────────────────────────────────────────────────
print("\nFetching preview...")
res = requests.get(f"{API}/preview/{gen_data['report_id']}")
assert res.status_code == 200, "Preview failed"
assert "Artificial Intelligence in Healthcare" in res.text
print("  Preview HTML contains topic title ✓")

# ── 5. Download DOCX ──────────────────────────────────────────────────────────
print("\nDownloading DOCX...")
res = requests.get(f"{API}/download-docx/{gen_data['report_id']}")
assert res.status_code == 200, f"DOCX download failed: {res.status_code}"
print(f"  DOCX size: {len(res.content)} bytes ✓")

print("\n✅ All tests passed!")

os.unlink(tmpfile.name)

"""Format-identical harness v2 — targets the EDITED reference
'PBL Report Rooftop Curtains.docx' (empty running headers, 12 media, 490 paras).

Runs the in-place pipeline with a NEW topic/team, then diffs output vs reference
on every format attribute that must stay identical, plus content-leakage checks.

Usage:  cd backend && ./venv/bin/python ../test_format_identical_v2.py
"""
import os
import re
import sys
import html
import zipfile

BACKEND = os.path.join(os.path.dirname(os.path.abspath(__file__)), "backend")
sys.path.insert(0, BACKEND)

# Load the Gemini key so the harness exercises LIVE content generation, matching
# the running backend. Without this, generate_content silently uses the offline
# fallback synthesizer (generic boilerplate) instead of real model output.
try:
    from dotenv import load_dotenv
    load_dotenv(os.path.join(BACKEND, ".env"))
except Exception:
    pass

from core.docx_reader import analyze_docx          # noqa: E402
from core.docx_writer import fill_template          # noqa: E402
from core.gemini_engine import generate_content     # noqa: E402

REFERENCE = "/Users/ajitreddy/Downloads/Test Reports/PBL Report Rooftop Curtains.docx"

TOPIC = "BEACHBUDDY INDIA"
STUDENTS = [
    ("AMEYA DESHPANDE", "S190234501"),
    ("SNEHA KULKARNI", "S190234502"),
    ("ROHIT JOSHI", "S190234503"),
    ("POOJA GAIKWAD", "S190234504"),
]
GROUP_NO = "Group No: 02"
YEAR = "2025-26"
GUIDE = "Prof. R. A. Vasmatkar"
HOD = "Dr. M. P. Wankhade"
PRINCIPAL = "Dr. S. D. Lokhande"

NOTES = ("Beachbuddy India is a beach-information website. A visitor opens the site, picks a beach, "
         "and sees live weather details — rain, air quality, humidity, wind — plus a recommendation on "
         "whether conditions are good to visit, swim, or surf. Built with HTML, CSS, TypeScript and JavaScript.")

# Old project markers that must NOT survive into the new report.
OLD_TOPIC_MARKERS = [
    "ROOFTOP CURTAIN", "TEMPERATURE SENSOR", "DS18B20", "ARDUINO UNO", "ARDUINO",
    "MOTOR DRIVER", "L23D", "OPEN-AIR CAFE", "MOTORIZED CURTAIN", "SOIL MOISTURE",
    "ONEWIRE", "DALLASTEMPERATURE",
]
# Friend's candidate + signature names (surnames/given names unique to the reference).
OLD_NAME_MARKERS = [
    "KOLPE", "KOLPUKE", "KONDE", "KSHIRSAGAR", "KUDALE",
    "PRATIKSHA", "VEDANT", "ATHRAV", "KRUSHNA", "TULSHIRAM", "SOMNATH", "GANESH",
]


def build_student_info():
    info = {
        "topic": TOPIC,
        "group_no": GROUP_NO,
        "academic_year": YEAR,
        "guide_name": GUIDE,
        "hod_name": HOD,
        "principal_name": PRINCIPAL,
        "college_name": "Sinhgad College of Engineering",
        "department_name": "Computer Engineering",
        "university_name": "Savitribai Phule Pune University",
        "report_type": "pbl",
    }
    for i, (name, seat) in enumerate(STUDENTS, 1):
        info[f"student_name_{i}"] = name
        info[f"roll_no_{i}"] = seat
    return info


def run_pipeline():
    doc_map_full = analyze_docx(REFERENCE)
    doc_map = {k: v for k, v in doc_map_full.items() if k != "file_path"}
    headings = doc_map_full.get("headings", [])
    student_info = build_student_info()

    print(f"[harness-v2] reference headings: {len(headings)}")
    ai_sections = generate_content(
        topic=TOPIC,
        sections=headings,
        student_info=student_info,
        reference_notes=NOTES,
        research_papers=[],
        source_material="",
        target_words_per_section=220,
    )
    print(f"[harness-v2] generated sections: {len(ai_sections)}")

    out = fill_template(
        source_docx_path=REFERENCE,
        doc_map=doc_map,
        student_info=student_info,
        ai_sections=ai_sections,
        images=[],
        output_filename="harness_beachbuddy.docx",
    )
    print(f"[harness-v2] wrote: {out}")
    return out, headings, ai_sections


# ---------- comparison helpers ----------

def _read_zip(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def _media_names(files):
    return sorted(n for n in files if n.startswith("word/media/"))


def _hf_parts(files):
    return sorted(n for n in files if re.match(r"word/(header|footer)\d*\.xml", n))


def _all_text(data):
    xml = data.decode("utf-8", "ignore")
    parts = re.findall(r"<w:t(?:\s[^>]*)?>(.*?)</w:t>", xml, re.S)
    return html.unescape("".join(parts))


def _sect_pr(path):
    files = _read_zip(path)
    xml = files["word/document.xml"].decode("utf-8", "ignore")
    pg = (re.search(r"<w:pgSz[^/]*/>", xml) or [None])
    pg = pg.group(0) if hasattr(pg, "group") else "(none)"
    m2 = re.search(r"<w:pgMar[^/]*/>", xml)
    mar = m2.group(0) if m2 else "(none)"
    return pg, mar, ("<w:titlePg" in xml)


def compare(ref, out, headings, ai_sections):
    results = []

    def check(name, ok, detail=""):
        results.append((name, ok, detail))

    rf, of = _read_zip(ref), _read_zip(out)

    rpg, rmar, rtp = _sect_pr(ref)
    opg, omar, otp = _sect_pr(out)
    check("page size identical", rpg == opg, f"ref={rpg} out={opg}")
    check("margins identical", rmar == omar, f"ref={rmar} out={omar}")
    check("titlePg flag unchanged", rtp == otp, f"ref={rtp} out={otp}")

    rm, om = _media_names(rf), _media_names(of)
    check("media files preserved", rm == om, f"ref={len(rm)} out={len(om)}")

    rh, oh = _hf_parts(rf), _hf_parts(of)
    check("header/footer parts preserved", rh == oh, f"ref={rh} out={oh}")

    # Headers in THIS reference are EMPTY — they must stay empty (no injected topic).
    ref_header_text = " ".join(_all_text(rf[n]) for n in rh if "header" in n).strip()
    out_header_text = " ".join(_all_text(of[n]) for n in oh if "header" in n).strip()
    if ref_header_text == "":
        check("headers preserved empty (no injected topic)", out_header_text == "",
              f"out_header={out_header_text[:120]!r}")
    else:
        check("header shows new topic", TOPIC.upper()[:15] in out_header_text.upper(),
              out_header_text[:120])
    check("header has no old topic", not any(m in out_header_text.upper() for m in OLD_TOPIC_MARKERS),
          out_header_text[:120])

    ref_footers = {n: re.sub(r"\s+", " ", _all_text(rf[n])).strip() for n in rh if "footer" in n}
    out_footers = {n: re.sub(r"\s+", " ", _all_text(of[n])).strip() for n in oh if "footer" in n}
    check("footers preserved verbatim", ref_footers == out_footers,
          f"ref={ref_footers} out={out_footers}")

    doc_text = _all_text(of["word/document.xml"]).upper()
    doc_norm = re.sub(r"\s+", " ", doc_text)
    check("new topic present in body/cover", TOPIC.upper()[:12] in doc_norm)
    for name, _ in STUDENTS:
        check(f"new student '{name}' present", name.upper() in doc_norm)
    leaked_topic = [m for m in OLD_TOPIC_MARKERS if m in doc_norm]
    check("no old topic leakage", not leaked_topic, f"found={leaked_topic}")
    leaked_names = [m for m in OLD_NAME_MARKERS if m in doc_norm]
    check("no old student-name leakage", not leaked_names, f"found={leaked_names}")

    def text_part(h):
        t = re.sub(r"\s+", " ", h.upper()).strip()
        return re.sub(r"^\d+(?:\.\d+)*\.?\s*", "", t).strip() or t

    missing = [h for h in headings if text_part(h) not in doc_norm]
    check("all reference headings preserved", not missing, f"missing={missing}")

    sample_keys = [h for h in ai_sections if re.sub(r"\s+", " ", h.upper()) in doc_norm][:3]
    check("generated sections landed in body", len(sample_keys) >= 1, f"keys={sample_keys}")

    # table structure preserved (List of Figures + TOC)
    from docx import Document
    rt, ot = Document(ref), Document(out)
    check("table count preserved", len(rt.tables) == len(ot.tables),
          f"ref={len(rt.tables)} out={len(ot.tables)}")
    if len(rt.tables) == len(ot.tables):
        shapes_ok = all((len(a.rows) == len(b.rows) and len(a.columns) == len(b.columns))
                        for a, b in zip(rt.tables, ot.tables))
        rshapes = [(len(a.rows), len(a.columns)) for a in rt.tables]
        oshapes = [(len(b.rows), len(b.columns)) for b in ot.tables]
        check("table shapes preserved", shapes_ok, f"ref={rshapes} out={oshapes}")

    return results


def main():
    out, headings, ai_sections = run_pipeline()
    results = compare(REFERENCE, out, headings, ai_sections)
    print("\n================ FORMAT / CONTENT DIFF (v2) ================")
    passed = 0
    for name, ok, detail in results:
        flag = "PASS" if ok else "FAIL"
        passed += ok
        line = f"[{flag}] {name}"
        if detail and not ok:
            line += f"  ->  {detail}"
        print(line)
    print("-----------------------------------------------------------")
    print(f"{passed}/{len(results)} checks passed")
    print(f"output: {out}")


if __name__ == "__main__":
    main()
